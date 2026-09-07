"""Excel 报告导出。"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

log = logging.getLogger("exporter")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import EXPORTS_DIR
from .database import get_db


def format_ms(ms: int) -> str:
    """毫秒 → 00:01:23.5（便于人读与视频播放器定位）。"""
    s, frac = divmod(int(ms), 1000)
    h, s = divmod(s, 3600)
    m, s = divmod(s, 60)
    return f"{h:02d}:{m:02d}:{s:02d}.{frac // 100}"


def export_hits(video_ids: list[int] | None = None) -> Path:
    """导出命中明细 Excel。video_ids 为空时导出全部。"""
    db = get_db()
    sql = """
        SELECT v.filename, v.path, h.word_text, c.name AS category,
               h.start_ms, h.end_ms, h.matched_text, h.sentence,
               v.transcribed_at
        FROM hits h
        JOIN videos v ON v.id = h.video_id
        JOIN words w ON w.id = h.word_id
        JOIN categories c ON c.id = w.category_id
    """
    params: tuple = ()
    if video_ids:
        sql += " WHERE h.video_id IN (%s)" % ",".join("?" * len(video_ids))
        params = tuple(video_ids)
    sql += " ORDER BY v.filename, h.start_ms"
    rows = db.query(sql, params)

    wb = Workbook()
    ws = wb.active
    ws.title = "违禁词检测报告"

    headers = ["视频文件", "分类", "违禁词", "命中原文", "开始时间", "结束时间",
               "完整句子", "视频路径", "检测时间"]
    header_fill = PatternFill("solid", fgColor="4F6EF7")
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([
            r["filename"], r["category"], r["word_text"], r["matched_text"],
            format_ms(r["start_ms"]), format_ms(r["end_ms"]),
            r["sentence"], r["path"], r["transcribed_at"] or "",
        ])

    widths = [28, 12, 12, 12, 11, 11, 50, 40, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = EXPORTS_DIR / f"违禁词检测报告_{stamp}.xlsx"
    wb.save(out)
    log.info("导出 Excel 报告: %s（%d 行命中）", out, len(rows))
    return out
